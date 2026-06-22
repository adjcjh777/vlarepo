from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
ASSET_DIR = ROOT / "assets" / "pattern_models_7_12"
OUT_DOCX = ROOT / "模式识别_第7-9模型教材小册子_程俊豪.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(39, 49, 58)
MUTED = RGBColor(95, 111, 122)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
PALE_YELLOW = "FFF3BF"
PALE_GREEN = "DDF4E8"


def set_run_font(run, size: float | None = None, bold: bool | None = None, color=None) -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths_in: list[float]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths_in):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def set_table_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "DADCE0")


def style_paragraph(paragraph, before=0, after=6, line=1.25, align=None) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if align is not None:
        paragraph.alignment = align


def add_para(doc, text: str = "", style: str | None = None, bold_prefix: str | None = None):
    p = doc.add_paragraph(style=style)
    style_paragraph(p)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, bold=True, color=INK)
        rest = p.add_run(text[len(bold_prefix):])
        set_run_font(rest, color=INK)
    else:
        r = p.add_run(text)
        set_run_font(r, color=INK)
    return p


def add_heading(doc, text: str, level: int) -> None:
    p = doc.add_heading(level=level)
    p.text = text
    style_paragraph(p, before=18 if level == 1 else 10, after=8 if level == 1 else 5)
    for run in p.runs:
        set_run_font(run, size={1: 16, 2: 13, 3: 12}.get(level, 11), bold=True, color=BLUE if level < 3 else DARK_BLUE)


def add_bullets(doc, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        style_paragraph(p, after=4, line=1.25)
        r = p.add_run(item)
        set_run_font(r, color=INK)


def add_numbers(doc, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        style_paragraph(p, after=4, line=1.25)
        r = p.add_run(item)
        set_run_font(r, color=INK)


def add_formula(doc, lines: list[str]) -> None:
    table = doc.add_table(rows=len(lines), cols=1)
    set_table_width(table, [6.3])
    set_table_borders(table)
    for row, line in zip(table.rows, lines):
        cell = row.cells[0]
        shade_cell(cell, LIGHT_GRAY)
        p = cell.paragraphs[0]
        style_paragraph(p, before=2, after=2, line=1.15, align=WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run(line)
        r.font.name = "Consolas"
        r._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
        r._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        r.font.size = Pt(10.5)
        r.font.color.rgb = DARK_BLUE


def add_callout(doc, label: str, text: str, fill: str = LIGHT_BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, [6.3])
    set_table_borders(table)
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    p = cell.paragraphs[0]
    style_paragraph(p, before=2, after=2, line=1.2)
    r1 = p.add_run(label + "：")
    set_run_font(r1, bold=True, color=DARK_BLUE)
    r2 = p.add_run(text)
    set_run_font(r2, color=INK)


def add_image(doc, rel_path: str, caption: str) -> None:
    path = ROOT / rel_path
    if not path.exists():
        return
    p = doc.add_paragraph()
    style_paragraph(p, before=6, after=3, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = p.add_run()
    run.add_picture(str(path), width=Cm(14.2))
    c = doc.add_paragraph()
    style_paragraph(c, before=0, after=8, line=1.1, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = c.add_run(caption)
    set_run_font(r, size=9.5, color=MUTED)


def add_simple_table(doc, headers: list[str], rows: list[list[str]], widths: list[float], header_fill: str = LIGHT_BLUE) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_width(table, widths)
    set_table_borders(table)
    for cell, header in zip(table.rows[0].cells, headers):
        shade_cell(cell, header_fill)
        p = cell.paragraphs[0]
        style_paragraph(p, after=0, line=1.15, align=WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run(header)
        set_run_font(r, bold=True, color=DARK_BLUE)
    for row_data in rows:
        row = table.add_row()
        for cell, value in zip(row.cells, row_data):
            p = cell.paragraphs[0]
            style_paragraph(p, after=0, line=1.15)
            r = p.add_run(value)
            set_run_font(r, size=10, color=INK)
    set_table_width(table, widths)


def setup_doc() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name in ("List Bullet", "List Number"):
        st = styles[style_name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(11)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.line_spacing = 1.25

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("模式识别教材小册子：感知机、GMM、HMM")
    set_run_font(r, size=9, color=MUTED)
    return doc


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    style_paragraph(p, before=30, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("模式识别经典模型教材小册子")
    set_run_font(r, size=24, bold=True, color=DARK_BLUE)
    p = doc.add_paragraph()
    style_paragraph(p, after=18, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("第 7-9 项：感知机、GMM、HMM")
    set_run_font(r, size=15, color=MUTED)
    add_callout(
        doc,
        "学习目标",
        "把读者当作第一次接触模式识别的小朋友：先用生活例子讲“它像什么”，再用很少的公式说明“机器怎么做”，最后用例题、习题和小实验巩固。",
        PALE_GREEN,
    )
    add_image(doc, "assets/pattern_models_7_12/overview.png", "图 0-1  第 7-12 项模型的三条主线，本小册子聚焦前三个模型。")
    add_simple_table(
        doc,
        ["模型", "一句话理解", "主要任务", "本章重点"],
        [
            ["感知机", "分错就修正的线性分类器", "二分类、线性边界学习", "权重更新与线性可分"],
            ["GMM", "多个高斯团块的概率混合", "软聚类、密度估计", "EM 算法与软归属"],
            ["HMM", "隐藏状态生成观测序列", "序列识别、路径解码", "三组概率与三类问题"],
        ],
        [1.35, 1.85, 1.55, 1.55],
    )
    doc.add_page_break()


def add_perceptron(doc: Document) -> None:
    add_heading(doc, "1. 感知机（Perceptron）", 1)
    add_callout(doc, "一句话定位", "感知机是一种线性二分类模型。它用一条线、一个平面或一个高维超平面把样本分成两类，并在分错样本时调整边界。")
    add_heading(doc, "1.1 核心直觉", 2)
    add_para(doc, "可以把感知机想象成一位刚学会分类的小老师。桌上有苹果和橘子，小老师拿一把尺子在桌面上画线：线左边算苹果，线右边算橘子。如果它把一个苹果分到橘子那边，就把尺子轻轻挪一下；如果把橘子分到苹果那边，也再挪一下。")
    add_para(doc, "所以，感知机最重要的想法很简单：先画一条分界线，错了就改，改到尽量少出错。二维数据中这条界线是一条直线；三维数据中是一个平面；维度更高时叫超平面。")
    add_image(doc, "assets/pattern_models_7_12/perceptron.png", "图 1-1  感知机根据误分类样本调整线性边界。")
    add_para(doc, "如果只记一句话，就是：感知机像一个会改错的分类小老师，分错一次就把线挪一点。")
    add_heading(doc, "1.2 必要公式", 2)
    add_formula(doc, ["s = w^T x + b", "f(x) = sign(s)", "w <- w + eta y_i x_i,    b <- b + eta y_i"])
    add_para(doc, "这些符号不用害怕。x 就是样本的特征，比如“颜色有多红、重量有多大”；w 表示每个特征有多重要；b 像一个整体门槛；eta 决定每次改错时挪多远。公式真正想说的是：机器先打分，再看正负号；如果分错，就按正确方向改一下。")
    add_heading(doc, "1.3 例题", 2)
    add_para(doc, "例题：小老师一开始完全不会分，设 w=(0,0)，b=0，学习率 eta=1。现在来了一个“苹果”样本 x=(2,1)，真实标签 y=+1。小老师分错了，应该怎么改？")
    add_callout(doc, "解答", "更新后 w=(0,0)+1*(+1)*(2,1)=(2,1)，b=0+1*(+1)=1。下一次再遇到这个样本时，s=2*2+1*1+1=6，模型会判为正类。", PALE_YELLOW)
    add_heading(doc, "1.4 运行流程", 2)
    add_numbers(doc, ["初始化权重和偏置。", "依次查看训练样本并计算当前预测。", "如果样本被分错，就按误差方向更新权重和偏置。", "多轮重复，直到错误很少或达到迭代上限。"])
    add_heading(doc, "1.5 习题", 2)
    add_bullets(doc, [
        "如果苹果和橘子混成一个圆圈套一个圆圈，一条直线还能分开吗？这能说明感知机的什么局限？",
        "如果每次改错时尺子挪得特别远，可能会发生什么？",
        "把“是否有促销词、是否有可疑链接、邮件长度”当作三个特征，解释每个权重像什么。",
    ])
    add_heading(doc, "1.6 小实验", 2)
    add_simple_table(
        doc,
        ["实验目标", "数据建议", "操作步骤", "观察指标"],
        [[
            "观察线性边界如何随错误样本移动",
            "二维人工数据，两类点大致可线性分开",
            "用 Python 生成散点；手写感知机更新；每轮画出边界",
            "错误数、边界位置、是否收敛",
        ]],
        [1.2, 1.55, 2.2, 1.4],
        PALE_GREEN,
    )
    add_heading(doc, "1.7 本节小结", 2)
    add_para(doc, "感知机的关键不是复杂公式，而是“分错就修正”。它把模式识别中的分类问题转化为寻找线性边界的问题，是理解神经网络和线性模型的起点。")


def add_gmm(doc: Document) -> None:
    add_heading(doc, "2. 高斯混合模型（GMM）", 1)
    add_callout(doc, "一句话定位", "GMM 认为数据来自多个高斯分布的叠加，每个样本可以用概率方式属于不同成分。")
    add_heading(doc, "2.1 核心直觉", 2)
    add_para(doc, "可以把 GMM 想象成一袋混在一起的糖果。糖果可能来自草莓味机器、葡萄味机器和柠檬味机器。我们现在只看到每颗糖的颜色和味道，不知道它到底是哪台机器做的。GMM 的任务，就是猜“每颗糖更像来自哪台机器”。")
    add_para(doc, "如果一颗糖颜色介于草莓和葡萄之间，GMM 不会粗暴地说它一定是草莓味，而可能说：它 60% 像草莓机器做的，35% 像葡萄机器做的，5% 像柠檬机器做的。这就是软归属。")
    add_image(doc, "assets/pattern_models_7_12/gmm.png", "图 2-1  GMM 用多个椭圆状高斯成分解释数据分布。")
    add_para(doc, "所以，GMM 的重点不是“强行分堆”，而是“带着不确定性分堆”。这很适合真实世界，因为很多东西本来就不是非黑即白。")
    add_heading(doc, "2.2 必要公式", 2)
    add_formula(doc, ["p(x) = sum_{k=1}^K pi_k N(x | mu_k, Sigma_k)", "gamma_{ik} = pi_k N(x_i | mu_k, Sigma_k) / sum_j pi_j N(x_i | mu_j, Sigma_j)"])
    add_para(doc, "公式里的每个高斯成分可以理解为一个“糖果制造机器”。mu_k 是这台机器最常做出的典型糖果，Sigma_k 描述糖果可能散开的范围和方向，pi_k 表示这台机器在整袋糖果里占多大比例。gamma_ik 叫责任度，可以理解为“这台机器对这颗糖有多大嫌疑”。")
    add_heading(doc, "2.3 EM 算法怎么理解", 2)
    add_numbers(doc, ["E 步：根据当前模型，估计每个样本属于每个成分的概率。", "M 步：根据这些概率，重新计算每个成分的位置、形状和权重。", "两步交替进行，模型逐渐贴合数据分布。"])
    add_heading(doc, "2.4 例题", 2)
    add_para(doc, "例题：一颗糖对草莓机器和葡萄机器的责任度分别是 0.75 和 0.25。如果只做硬分类，它会被分到哪一类？如果做软聚类，这两个数字又告诉我们什么？")
    add_callout(doc, "解答", "硬分类会把它分到草莓机器。但软聚类还告诉我们：它不是百分百草莓，也有 25% 像葡萄。这种“不确定性”有时比单纯给一个类别更有用。", PALE_YELLOW)
    add_heading(doc, "2.5 习题", 2)
    add_bullets(doc, [
        "如果糖果有三种来源，却只让 GMM 找两个团块，会发生什么？",
        "如果让 GMM 找太多团块，它会不会把一些偶然的小差异也当成一种来源？",
        "如果一颗糖不像任何一台机器做出来的，它为什么可能是异常样本？",
    ])
    add_heading(doc, "2.6 小实验", 2)
    add_simple_table(
        doc,
        ["实验目标", "数据建议", "操作步骤", "观察指标"],
        [[
            "比较硬聚类与软聚类的差别",
            "二维混合高斯数据，三个不同方向的簇",
            "分别运行 K-Means 与 GaussianMixture；画出簇中心、椭圆和责任度",
            "AIC/BIC、责任度、边界样本归属",
        ]],
        [1.2, 1.55, 2.2, 1.4],
        PALE_GREEN,
    )
    add_heading(doc, "2.7 本节小结", 2)
    add_para(doc, "GMM 的核心价值是用概率方式描述“混合”。它不仅告诉我们样本属于哪个群体，还告诉我们这种判断有多确定。")


def add_hmm(doc: Document) -> None:
    add_heading(doc, "3. 隐马尔可夫模型（HMM）", 1)
    add_callout(doc, "一句话定位", "HMM 用于序列数据。它假设背后存在一串看不见的状态，这些状态按时间变化，并生成我们能观察到的信号。")
    add_heading(doc, "3.1 核心直觉", 2)
    add_para(doc, "可以把 HMM 想象成“看脚印猜路线”的游戏。一个小朋友在几个房间里走来走去，但我们没有看到他本人，只看到了地上的脚印。房间就是隐藏状态，脚印就是观测。我们要根据一串脚印，猜他最可能走过哪些房间。")
    add_para(doc, "很多任务都像这个游戏：语音识别时，我们看见的是声音波形，却想知道背后的音素或词；天气例子里，我们看见某人有没有带伞，却想猜当天是晴天、阴天还是雨天。")
    add_image(doc, "assets/pattern_models_7_12/hmm.png", "图 3-1  HMM 的隐藏状态链与观测序列。")
    add_para(doc, "所以，HMM 的重点是：看见一串表面现象，推测背后一串看不见的状态。")
    add_heading(doc, "3.2 三组关键概率", 2)
    add_formula(doc, ["pi_i = P(q_1 = i)", "a_ij = P(q_{t+1} = j | q_t = i)", "b_j(o) = P(o_t = o | q_t = j)"])
    add_para(doc, "用脚印故事来理解：初始概率表示小朋友一开始在哪个房间的可能性；转移概率表示他从一个房间走到另一个房间的可能性；发射概率表示某个房间留下某种脚印的可能性。")
    add_heading(doc, "3.3 三个经典问题", 2)
    add_simple_table(
        doc,
        ["问题", "在问什么", "典型算法"],
        [
            ["评估", "给定模型和观测序列，这段观测出现的概率是多少？", "前向算法"],
            ["解码", "给定观测序列，最可能的隐藏状态路径是什么？", "Viterbi 算法"],
            ["学习", "只有观测序列时，如何估计模型参数？", "Baum-Welch 算法"],
        ],
        [1.1, 3.7, 1.5],
    )
    add_heading(doc, "3.4 例题", 2)
    add_para(doc, "例题：在“看脚印猜路线”中，我们看到的序列是：湿脚印、湿脚印、干脚印。请问隐藏状态和观测分别是什么？")
    add_callout(doc, "解答", "观测是我们真正看到的脚印类型，比如湿脚印、干脚印；隐藏状态是小朋友走过的房间，比如水池房、走廊、教室。HMM 想做的事，就是根据脚印序列猜最可能的房间路线。", PALE_YELLOW)
    add_heading(doc, "3.5 习题", 2)
    add_bullets(doc, [
        "在天气和带伞例子里，哪个是隐藏状态？哪个是观测？",
        "为什么只看今天有没有带伞，可能还不如看连续三天有没有带伞？",
        "如果小朋友可能记得很久以前走过哪里，只看上一个房间会不会太简单？这说明 HMM 有什么局限？",
    ])
    add_heading(doc, "3.6 小实验", 2)
    add_simple_table(
        doc,
        ["实验目标", "数据建议", "操作步骤", "观察指标"],
        [[
            "用 Viterbi 解码最可能状态路径",
            "手工构造天气状态和带伞观测序列",
            "设置初始、转移、发射概率；输入观测序列；输出最可能天气序列",
            "路径概率、预测状态、与直觉是否一致",
        ]],
        [1.2, 1.55, 2.2, 1.4],
        PALE_GREEN,
    )
    add_heading(doc, "3.7 本节小结", 2)
    add_para(doc, "HMM 把序列识别拆成“隐藏状态”和“可见观测”。看懂这两层结构，就能理解它为什么适合处理语音、文本、动作和故障这类随时间展开的问题。")


def add_comparison_and_refs(doc: Document) -> None:
    add_heading(doc, "4. 三个模型的横向比较", 1)
    add_simple_table(
        doc,
        ["模型", "类型", "核心思想", "适合问题", "主要局限"],
        [
            ["感知机", "判别式线性分类", "分错就修正边界", "线性二分类、神经网络入门", "只能表达线性边界"],
            ["GMM", "生成式概率模型", "多个概率团块混合产生数据", "软聚类、密度估计、异常检测", "需指定成分数，对初始化敏感"],
            ["HMM", "序列概率模型", "隐藏状态按时间生成观测", "语音、文本、动作、故障序列", "难表达长期依赖"],
        ],
        [1.0, 1.2, 1.7, 1.45, 1.35],
    )
    add_para(doc, "这三个模型分别代表三种思路：感知机关注分类边界，GMM 关注概率分布的混合，HMM 关注序列背后的状态变化。它们已经不是所有工业任务中的最前沿工具，但仍然是理解现代机器学习的重要地基。")
    add_heading(doc, "5. 综合复习题", 1)
    add_numbers(doc, [
        "比较感知机与 GMM：一个为什么叫判别式模型，另一个为什么叫生成式模型？",
        "如果一个任务既有类别标签，又希望解释数据来源的不确定性，应该优先考虑哪类思想？",
        "请举一个 HMM 比普通分类器更合适的任务，并说明原因。",
        "任选一个模型，说明它在现代深度学习系统中仍然保留的思想价值。",
    ])
    add_heading(doc, "6. 参考文献与延伸阅读", 1)
    add_numbers(doc, [
        "Richard O. Duda, Peter E. Hart, David G. Stork. Pattern Classification.",
        "Christopher M. Bishop. Pattern Recognition and Machine Learning.",
        "Tom M. Mitchell. Machine Learning.",
        "Kevin P. Murphy. Machine Learning: A Probabilistic Perspective.",
        "Lawrence R. Rabiner. A Tutorial on Hidden Markov Models and Selected Applications in Speech Recognition.",
        "scikit-learn User Guide: Perceptron, Gaussian Mixture Models, and model evaluation examples.",
    ])


def main() -> None:
    doc = setup_doc()
    add_cover(doc)
    add_perceptron(doc)
    add_gmm(doc)
    add_hmm(doc)
    add_comparison_and_refs(doc)
    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    main()
