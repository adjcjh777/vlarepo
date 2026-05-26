from __future__ import annotations

import math
import shutil
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parent
ASSET_DIR = ROOT / "assets" / "pattern_models_7_12"
MD_PATH = ROOT / "模式识别_7-12模型研究与教学资料.md"
DOCX_PATH = ROOT / "模式识别_7-12模型研究与教学资料.docx"

FONT_CANDIDATES = [
    "/System/Library/Fonts/STHeiti Medium.ttc",
    "/System/Library/Fonts/PingFang.ttc",
    "/Library/Fonts/Arial Unicode.ttf",
]

COLORS = {
    "ink": "#27313a",
    "muted": "#5f6f7a",
    "line": "#9aa8b2",
    "blue": "#2f80ed",
    "blue_light": "#d7e9ff",
    "red": "#e45757",
    "red_light": "#ffe1e1",
    "green": "#2f9e6d",
    "green_light": "#ddf4e8",
    "yellow": "#f2c94c",
    "yellow_light": "#fff3bf",
    "purple": "#8b5cf6",
    "purple_light": "#eee7ff",
    "bg": "#fbfcfe",
}


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    for candidate in FONT_CANDIDATES:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size=size, index=0)
    return ImageFont.load_default()


def text_size(draw: ImageDraw.ImageDraw, text: str, fnt: ImageFont.ImageFont) -> tuple[int, int]:
    box = draw.textbbox((0, 0), text, font=fnt)
    return box[2] - box[0], box[3] - box[1]


def draw_centered(
    draw: ImageDraw.ImageDraw,
    xy: tuple[float, float],
    text: str,
    fnt: ImageFont.ImageFont,
    fill: str = COLORS["ink"],
) -> None:
    w, h = text_size(draw, text, fnt)
    draw.text((xy[0] - w / 2, xy[1] - h / 2), text, font=fnt, fill=fill)


def rounded_label(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    title: str,
    subtitle: str,
    fill: str,
    outline: str,
) -> None:
    draw.rounded_rectangle(box, radius=22, fill=fill, outline=outline, width=3)
    x1, y1, x2, y2 = box
    draw_centered(draw, ((x1 + x2) / 2, y1 + 35), title, font(28), COLORS["ink"])
    draw_centered(draw, ((x1 + x2) / 2, y1 + 78), subtitle, font(20), COLORS["muted"])


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], fill: str = COLORS["line"], width: int = 4) -> None:
    draw.line([start, end], fill=fill, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    length = 16
    for delta in (math.pi * 0.82, -math.pi * 0.82):
        x = end[0] + length * math.cos(angle + delta)
        y = end[1] + length * math.sin(angle + delta)
        draw.line([end, (x, y)], fill=fill, width=width)


def base_canvas(width: int = 1600, height: int = 900, title: str | None = None) -> tuple[Image.Image, ImageDraw.ImageDraw]:
    img = Image.new("RGB", (width, height), COLORS["bg"])
    draw = ImageDraw.Draw(img)
    if title:
        draw.text((70, 48), title, font=font(42), fill=COLORS["ink"])
        draw.line((70, 112, width - 70, 112), fill="#d9e2ec", width=3)
    return img, draw


def save(img: Image.Image, name: str) -> str:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    path = ASSET_DIR / name
    img.save(path, "PNG")
    return str(path.relative_to(ROOT))


def use_codex_generated(name: str) -> str:
    source = ASSET_DIR / "codex_generated" / name
    target = ASSET_DIR / name
    if not source.exists():
        raise FileNotFoundError(f"Missing Codex generated image: {source}")
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    im = Image.open(source).convert("RGB")
    w, h = im.size
    target_ratio = 16 / 9
    if w / h > target_ratio:
        new_w = int(h * target_ratio)
        left = (w - new_w) // 2
        im = im.crop((left, 0, left + new_w, h))
    elif w / h < target_ratio:
        new_h = int(w / target_ratio)
        top = (h - new_h) // 2
        im = im.crop((0, top, w, top + new_h))
    im = im.resize((1600, 900), Image.Resampling.LANCZOS)
    im.save(target, "PNG", optimize=True)
    return str(target.relative_to(ROOT))


def draw_overview() -> str:
    img, draw = base_canvas(title="第 7-12 项模型的三条主线")
    lanes = [
        ("判别式分类", "直接学习分类边界", COLORS["blue_light"], COLORS["blue"], ["感知机", "支持向量机"]),
        ("概率生成式建模", "解释数据如何产生", COLORS["green_light"], COLORS["green"], ["GMM", "HMM", "贝叶斯网络"]),
        ("实例学习", "根据相似样本判断", COLORS["yellow_light"], "#d59a00", ["K-近邻"]),
    ]
    x_positions = [115, 600, 1085]
    for x, (title, sub, fill, outline, models) in zip(x_positions, lanes):
        draw.rounded_rectangle((x, 175, x + 400, 675), radius=28, fill=fill, outline=outline, width=4)
        draw_centered(draw, (x + 200, 225), title, font(32), COLORS["ink"])
        draw_centered(draw, (x + 200, 270), sub, font(22), COLORS["muted"])
        for i, model in enumerate(models):
            y = 355 + i * 90
            draw.rounded_rectangle((x + 70, y, x + 330, y + 56), radius=18, fill="white", outline=outline, width=3)
            draw_centered(draw, (x + 200, y + 28), model, font(26), COLORS["ink"])
    arrow(draw, (515, 425), (600, 425), "#b8c3cc", 5)
    arrow(draw, (1000, 425), (1085, 425), "#b8c3cc", 5)
    draw.text((135, 745), "同一个识别任务可以用不同思想解决：边界、概率、相似性分别提供不同视角。", font=font(28), fill=COLORS["ink"])
    return save(img, "overview.png")


def draw_perceptron() -> str:
    img, draw = base_canvas(title="感知机：由错误样本推动线性边界")
    draw.line((190, 760, 1390, 760), fill=COLORS["line"], width=4)
    draw.line((190, 760, 190, 170), fill=COLORS["line"], width=4)
    blue_points = [(330, 610), (420, 570), (500, 650), (560, 540), (630, 590), (720, 515)]
    red_points = [(800, 410), (890, 330), (960, 390), (1040, 300), (1130, 365), (1210, 260)]
    mistake = (725, 440)
    for p in blue_points:
        draw.ellipse((p[0] - 14, p[1] - 14, p[0] + 14, p[1] + 14), fill=COLORS["blue"], outline="white", width=3)
    for p in red_points:
        draw.ellipse((p[0] - 14, p[1] - 14, p[0] + 14, p[1] + 14), fill=COLORS["red"], outline="white", width=3)
    draw.ellipse((mistake[0] - 17, mistake[1] - 17, mistake[0] + 17, mistake[1] + 17), fill=COLORS["red"], outline=COLORS["ink"], width=4)
    draw.line((430, 705, 1110, 190), fill=COLORS["ink"], width=5)
    draw.line((470, 735, 1150, 220), fill="#c5ced6", width=3)
    draw.line((390, 675, 1070, 160), fill="#c5ced6", width=3)
    arrow(draw, (690, 385), (755, 445), COLORS["purple"], 5)
    draw.text((800, 470), "误分类点带来一次参数更新", font=font(26), fill=COLORS["purple"])
    draw.text((1010, 710), "直线两侧代表两个类别", font=font(24), fill=COLORS["muted"])
    return save(img, "perceptron.png")


def draw_gmm() -> str:
    img, draw = base_canvas(title="GMM：一个样本可以按概率属于多个高斯成分")
    img = img.convert("RGBA")
    clusters = [
        ((500, 470), (420, 210), -18, COLORS["blue_light"], COLORS["blue"]),
        ((855, 405), (500, 250), 18, COLORS["red_light"], COLORS["red"]),
        ((910, 610), (420, 200), -10, COLORS["green_light"], COLORS["green"]),
    ]
    for center, size, angle, fill, outline in clusters:
        layer = Image.new("RGBA", img.size, (0, 0, 0, 0))
        d = ImageDraw.Draw(layer)
        cx, cy = center
        w, h = size
        d.ellipse((cx - w / 2, cy - h / 2, cx + w / 2, cy + h / 2), fill=fill + "b8", outline=outline, width=5)
        layer = layer.rotate(angle, center=center, resample=Image.Resampling.BICUBIC)
        img.alpha_composite(layer)
    draw = ImageDraw.Draw(img)
    for cx, cy, color in [
        (460, 470, COLORS["blue"]), (520, 420, COLORS["blue"]), (570, 520, COLORS["blue"]),
        (835, 375, COLORS["red"]), (930, 435, COLORS["red"]), (760, 430, COLORS["red"]),
        (910, 610, COLORS["green"]), (1000, 650, COLORS["green"]), (830, 660, COLORS["green"]),
    ]:
        draw.ellipse((cx - 12, cy - 12, cx + 12, cy + 12), fill=color, outline="white", width=3)
    draw.ellipse((720, 505, 760, 545), fill=COLORS["yellow"], outline=COLORS["ink"], width=4)
    draw.text((790, 505), "边界附近的样本：\n不是非黑即白，而是有多个归属概率", font=font(26), fill=COLORS["ink"], spacing=8)
    draw.text((260, 725), "每个椭圆表示一个高斯成分；椭圆方向和大小由协方差决定。", font=font(26), fill=COLORS["muted"])
    return save(img.convert("RGB"), "gmm.png")


def draw_hmm() -> str:
    img, draw = base_canvas(title="HMM：隐藏状态链生成可见观测")
    hidden = [(360, 285, "天气1"), (710, 285, "天气2"), (1060, 285, "天气3")]
    obs = [(360, 585, "是否带伞1"), (710, 585, "是否带伞2"), (1060, 585, "是否带伞3")]
    for i in range(len(hidden) - 1):
        arrow(draw, (hidden[i][0] + 95, hidden[i][1]), (hidden[i + 1][0] - 95, hidden[i + 1][1]), COLORS["green"], 5)
    for (hx, hy, _), (ox, oy, _) in zip(hidden, obs):
        arrow(draw, (hx, hy + 55), (ox, oy - 65), COLORS["purple"], 5)
    for x, y, label in hidden:
        draw.rounded_rectangle((x - 105, y - 55, x + 105, y + 55), radius=22, fill=COLORS["green_light"], outline=COLORS["green"], width=4)
        draw_centered(draw, (x, y), label, font(28))
    for x, y, label in obs:
        draw.rounded_rectangle((x - 125, y - 58, x + 125, y + 58), radius=22, fill=COLORS["purple_light"], outline=COLORS["purple"], width=4)
        draw_centered(draw, (x, y), label, font(26))
    draw.text((260, 720), "上层状态看不见，只能通过下层观测反推；相邻状态之间按转移概率变化。", font=font(28), fill=COLORS["ink"])
    return save(img, "hmm.png")


def draw_bayes() -> str:
    img, draw = base_canvas(title="贝叶斯网络：把变量依赖画成概率图")
    nodes = {
        "感染": (490, 260, COLORS["red_light"], COLORS["red"]),
        "过敏": (870, 260, COLORS["yellow_light"], "#d59a00"),
        "发烧": (335, 505, COLORS["red_light"], COLORS["red"]),
        "咳嗽": (680, 505, COLORS["green_light"], COLORS["green"]),
        "检测阳性": (1035, 505, COLORS["blue_light"], COLORS["blue"]),
    }
    edges = [("感染", "发烧"), ("感染", "咳嗽"), ("感染", "检测阳性"), ("过敏", "咳嗽")]
    for a, b in edges:
        ax, ay, _, _ = nodes[a]
        bx, by, _, _ = nodes[b]
        arrow(draw, (ax, ay + 60), (bx, by - 70), COLORS["line"], 5)
    for label, (x, y, fill, outline) in nodes.items():
        draw.rounded_rectangle((x - 110, y - 60, x + 110, y + 60), radius=24, fill=fill, outline=outline, width=4)
        draw_centered(draw, (x, y), label, font(28))
    draw.text((230, 705), "边表示直接依赖；没有边不等于毫无关系，而是关系可由其他变量解释。", font=font(28), fill=COLORS["ink"])
    return save(img, "bayes.png")


def draw_knn() -> str:
    img, draw = base_canvas(title="KNN：由附近样本投票决定类别")
    draw.line((215, 760, 1390, 760), fill=COLORS["line"], width=4)
    draw.line((215, 760, 215, 170), fill=COLORS["line"], width=4)
    blue_points = [(350, 560), (410, 500), (470, 610), (560, 545), (620, 640), (710, 590)]
    red_points = [(815, 445), (900, 390), (970, 510), (1065, 430), (1160, 520), (1215, 345)]
    query = (760, 515)
    for p in blue_points:
        draw.rectangle((p[0] - 13, p[1] - 13, p[0] + 13, p[1] + 13), fill=COLORS["blue"], outline="white", width=3)
    for p in red_points:
        draw.ellipse((p[0] - 14, p[1] - 14, p[0] + 14, p[1] + 14), fill=COLORS["red"], outline="white", width=3)
    draw.ellipse((query[0] - 155, query[1] - 155, query[0] + 155, query[1] + 155), outline=COLORS["purple"], width=5)
    draw.regular_polygon((query, 30), n_sides=5, rotation=-90, fill=COLORS["yellow"], outline=COLORS["ink"])
    draw.text((950, 655), "K=5 时看最近的 5 个邻居", font=font(28), fill=COLORS["purple"])
    draw.text((295, 710), "方块/圆点是已知样本；星形是待分类样本。", font=font(24), fill=COLORS["muted"])
    return save(img, "knn.png")


def draw_svm() -> str:
    img, draw = base_canvas(title="SVM：选择间隔最大的分类边界")
    draw.line((205, 760, 1395, 760), fill=COLORS["line"], width=4)
    draw.line((205, 760, 205, 165), fill=COLORS["line"], width=4)
    boundary = ((485, 720), (1075, 180))
    margin1 = ((410, 650), (1000, 110))
    margin2 = ((570, 790), (1160, 250))
    draw.line(margin1, fill="#c5ced6", width=4)
    draw.line(margin2, fill="#c5ced6", width=4)
    draw.line(boundary, fill=COLORS["ink"], width=6)
    blue_points = [(380, 600), (455, 530), (520, 610), (610, 520), (700, 455)]
    red_points = [(795, 365), (875, 290), (930, 420), (1035, 330), (1120, 250)]
    support = [(700, 455), (795, 365), (610, 520)]
    for p in blue_points:
        draw.ellipse((p[0] - 15, p[1] - 15, p[0] + 15, p[1] + 15), fill=COLORS["blue"], outline="white", width=3)
    for p in red_points:
        draw.ellipse((p[0] - 15, p[1] - 15, p[0] + 15, p[1] + 15), fill=COLORS["red"], outline="white", width=3)
    for p in support:
        draw.ellipse((p[0] - 27, p[1] - 27, p[0] + 27, p[1] + 27), outline=COLORS["purple"], width=5)
    draw.text((1030, 575), "被圈出的点决定边界，\n它们就是支持向量。", font=font(28), fill=COLORS["purple"], spacing=8)
    draw.text((300, 720), "中线是分类边界，两侧浅线表示间隔。", font=font(24), fill=COLORS["muted"])
    return save(img, "svm.png")


def build_markdown(paths: dict[str, str]) -> str:
    return f"""# 模式识别作业三：第 7-12 个模型研究与教学资料

作者：程俊豪  
协作队友：傅楷  
选题范围：根据《现代模式识别的主流模型清单》，本组完成清单第 7-12 项模型：感知机、高斯混合模型、隐马尔可夫模型、贝叶斯网络、K-近邻、支持向量机。  
修订定位：面向 AI 兴趣爱好者和课程初学者，正文以“能读懂、能举例、能比较”为目标。公式只保留理解模型所必需的骨架，更多篇幅用于解释公式背后的问题。

## 0. 六个模型的共同主线

模式识别的核心问题，是让机器从数据中识别类别、结构或规律。第 7-12 个模型刚好覆盖三种常见思路：直接学习分类边界、用概率解释数据生成过程、根据相似样本做判断。

![图 0-1：第 7-12 项模型的三条主线]({paths["overview"]}){{width=92%}}

感知机和支持向量机都关心“边界”。感知机更像一个会从错误中改正的线性分类器；支持向量机则进一步追求最稳的边界。高斯混合模型、隐马尔可夫模型和贝叶斯网络都关心“概率”。它们不只判断类别，还试图解释样本、序列或变量关系是如何产生的。K-近邻的思路最接近日常经验：新样本和谁最像，就更可能属于哪一类。

把六个模型放在一起看，可以得到一张清晰地图：感知机是神经网络的入口，GMM 是软聚类和密度估计的代表，HMM 是经典序列建模工具，贝叶斯网络是概率推理框架，KNN 是实例学习基线，SVM 是最大间隔分类器。它们有些已经不再是最前沿的工业主力，但仍然是理解现代机器学习的重要地基。

## 清单第 7 项：感知机（Perceptron）

### 1.1 一句话定位

感知机是一种线性二分类模型。它用一条线、一个平面或一个高维超平面把样本分成两类，并在分错样本时调整边界。

### 1.2 核心直觉

可以把感知机想象成一把会移动的尺子。二维数据中，尺子对应一条直线；三维数据中，对应一个平面；维度更高时，对应超平面。模型的目标不是记住每个样本，而是找到一条能把两类样本分开的界线。

![图 1-1：感知机根据误分类样本调整线性边界]({paths["perceptron"]}){{width=92%}}

感知机的学习方式很朴素：如果一个正类样本被分到负类一侧，边界就向有利于正类的方向调整；如果负类样本被分错，边界就向相反方向调整。每次调整都只利用当前这个错误样本，因此感知机也体现了在线学习的思想。

### 1.3 一个容易理解的例子

假设要根据两个特征判断邮件是否为垃圾邮件：第一个特征是“是否包含促销词”，第二个特征是“是否包含可疑链接”。如果一封正常邮件被误判为垃圾邮件，模型会降低某些特征的权重；如果一封垃圾邮件被漏掉，模型会提高相关特征的权重。经过多轮修正后，模型形成一条区分两类邮件的直线边界。

这个例子说明，权重不是抽象数字，而是模型对每个特征重要性的判断。偏置项则像整体门槛，决定模型倾向于更宽松还是更严格。

### 1.4 必要公式

感知机先计算特征的加权和：

$$
s=w^Tx+b
$$

其中 $w$ 是权重，$x$ 是样本特征，$b$ 是偏置。然后根据符号得到类别：

$$
f(x)=sign(s)
$$

当样本 $(x_i,y_i)$ 被分错时，参数更新为：

$$
w \\leftarrow w+\\eta y_i x_i, \\quad b \\leftarrow b+\\eta y_i
$$

这条公式的含义并不复杂：$y_i x_i$ 决定调整方向，学习率 $\\eta$ 决定每次移动多远。

### 1.5 运行流程

1. 初始化权重和偏置。
2. 依次查看训练样本。
3. 用当前边界判断样本类别。
4. 如果分错，就按误差方向更新权重和偏置。
5. 多轮重复，直到错误很少或达到迭代上限。

如果数据确实能被一条线分开，感知机最终可以找到一个可分边界。若数据本身线性不可分，模型会反复调整而难以稳定。

### 1.6 适用场景与局限

感知机适合用来理解线性分类器、权重更新和神经网络基本单元。它结构简单，计算成本低，也能作为更复杂模型的入门版本。

它的主要局限是只能表达线性边界。对于 XOR 这类必须用弯曲边界才能区分的问题，单层感知机无能为力。它也不直接输出概率，对样本顺序和学习率比较敏感。

### 1.7 本节小结

感知机的关键不是复杂公式，而是“分错就修正”。它把模式识别中的分类问题转化为寻找线性边界的问题，是理解神经网络和线性模型的起点。

## 清单第 8 项：高斯混合模型（GMM）

### 2.1 一句话定位

高斯混合模型认为数据来自多个高斯分布的叠加，每个样本可以用概率方式属于不同成分。

### 2.2 核心直觉

现实数据经常不是单一人群或单一来源产生的。一个班级的身高、一个城市的消费水平、一批传感器数据的正常与异常状态，都可能由多个子群体混合而成。GMM 的想法是：不要强行把数据看成一个整体，而是拆成若干个“概率团块”。

![图 2-1：GMM 用多个椭圆状高斯成分解释数据分布]({paths["gmm"]}){{width=92%}}

与 K-Means 不同，GMM 不要求一个样本只能属于一个簇。边界附近的样本可以 60% 属于第一个成分、35% 属于第二个成分、5% 属于第三个成分。这种“软归属”让模型更贴近真实世界的不确定性。

### 2.3 一个容易理解的例子

假设商场记录了顾客的“年消费金额”和“到店频率”。直接看散点图时，顾客群体可能混在一起。GMM 可以把他们分成几个概率群体：高频高消费的核心顾客、低频高消费的偶发大额顾客、低频低消费的普通顾客。一个位于两个群体中间的顾客，不会被硬塞进某一类，而是得到一组归属概率。

这种结果比单纯分类更丰富，因为它既给出分组，也给出不确定程度。

### 2.4 必要公式

GMM 的整体概率密度写作：

$$
p(x)=\\sum_{{k=1}}^K \\pi_k N(x|\\mu_k,\\Sigma_k)
$$

这里的 $K$ 是高斯成分数，$\\pi_k$ 是第 $k$ 个成分的权重，$\\mu_k$ 表示中心位置，$\\Sigma_k$ 表示形状和方向。公式表达的是：一个样本的概率由多个高斯成分加权相加得到。

### 2.5 EM 算法怎么理解

GMM 通常用 EM 算法学习参数。EM 可以理解为反复做两件事：

1. E 步：根据当前模型，估计每个样本属于每个成分的概率。
2. M 步：根据这些概率，重新计算每个成分的位置、形状和权重。

E 步像是在重新分配样本责任，M 步像是在根据责任重新摆放每个椭圆。两步交替进行，模型会逐渐贴合数据分布。

### 2.6 适用场景与局限

GMM 常用于软聚类、密度估计、异常检测、语音建模和背景建模。它比 K-Means 更灵活，因为每个簇可以有不同大小、方向和不确定性。

局限也很明显：成分数 $K$ 需要提前设定；初始化不好会影响结果；高维数据中协方差矩阵难以稳定估计；如果真实数据分布完全不像高斯混合，模型解释力会下降。

### 2.7 本节小结

GMM 的核心价值是用概率方式描述“混合”。它不仅告诉我们样本属于哪个群体，还告诉我们这种判断有多确定。

## 清单第 9 项：隐马尔可夫模型（HMM）

### 3.1 一句话定位

隐马尔可夫模型用于序列数据。它假设背后存在一串看不见的状态，这些状态按时间变化，并生成我们能观察到的信号。

### 3.2 核心直觉

许多模式识别任务不是孤立样本，而是连续过程。语音是一串声音片段，动作是一串姿态变化，文本是一串词语。HMM 用两层结构描述这类问题：上层是隐藏状态，下层是观测结果。

![图 3-1：HMM 的隐藏状态链与观测序列]({paths["hmm"]}){{width=92%}}

天气与带伞是最常见的例子。真实天气可能是晴天、阴天或雨天，但我们只看到某人是否带伞。天气是隐藏状态，带伞行为是观测。HMM 的任务是根据观测序列，推断背后的状态变化，或者计算某段观测出现的概率。

### 3.3 一个容易理解的例子

在语音识别中，真正想识别的是词或音素，但机器直接拿到的是声波特征。一个词会生成连续的声音片段，声音片段之间又有时间顺序。HMM 可以把“音素状态”放在隐藏层，把“声学特征”放在观测层，从而把语音识别变成序列推断问题。

在深度学习普及之前，HMM 长期是语音识别的重要基础。即使今天许多系统已经改用神经网络，HMM 的状态转移思想仍然很适合理解序列建模。

### 3.4 三组关键概率

HMM 通常由三类概率描述：

$$
\\pi_i=P(q_1=i)
$$

初始概率表示序列一开始处于某个状态的可能性。

$$
a_{{ij}}=P(q_{{t+1}}=j|q_t=i)
$$

转移概率表示当前状态如何变成下一个状态。

$$
b_j(o)=P(o_t=o|q_t=j)
$$

发射概率表示某个隐藏状态生成某个观测的可能性。

### 3.5 三个经典问题

HMM 经常处理三类问题：

1. 评估：给定模型和观测序列，计算这段观测出现的概率。
2. 解码：给定模型和观测序列，找出最可能的隐藏状态路径。
3. 学习：给定观测序列，估计初始概率、转移概率和发射概率。

对应算法分别包括前向算法、Viterbi 算法和 Baum-Welch 算法。初学时不必先背算法细节，先理解三类问题分别在问什么更重要。

### 3.6 适用场景与局限

HMM 适合语音识别、词性标注、手势识别、生物序列分析和故障诊断等序列任务。它的优势是结构清晰、可解释性强，能明确表达状态随时间转移的过程。

它的局限来自假设较强：当前状态通常只依赖上一个状态，观测也被简化为由当前状态生成。现实任务中的长期依赖更复杂，因此许多现代序列任务会使用 RNN、LSTM 或 Transformer。

### 3.7 本节小结

HMM 把序列识别拆成“隐藏状态”和“可见观测”。看懂这两层结构，就能理解它为什么适合处理语音、文本、动作和故障这类随时间展开的问题。

## 清单第 10 项：贝叶斯网络（Bayesian Network）

### 4.1 一句话定位

贝叶斯网络用有向无环图表示变量之间的条件依赖关系，并用概率完成不确定推理。

### 4.2 核心直觉

在复杂问题里，变量之间很少完全独立。疾病可能导致发烧和咳嗽，过敏也可能导致咳嗽，检测阳性又会改变我们对疾病的判断。贝叶斯网络把这些变量画成图，再给每个节点配上条件概率。

![图 4-1：贝叶斯网络用有向图组织变量依赖]({paths["bayes"]}){{width=92%}}

图结构的价值在于把复杂问题拆开。完整联合概率表会非常庞大，而贝叶斯网络只需要描述每个变量受哪些直接父节点影响。这样既减少参数，也让模型结构更容易解释。

### 4.3 一个容易理解的例子

在医疗诊断中，已知“发烧”和“检测阳性”会提高感染概率；如果同时发现“过敏”，咳嗽对感染的支持力度可能下降。这种现象叫解释消除：一个症状可能由多个原因解释，当其中一个原因被确认时，另一个原因的概率会发生变化。

这正是贝叶斯网络擅长的地方。它不是只给出一个类别，而是在证据变化时不断更新概率判断。

### 4.4 必要公式

贝叶斯网络的联合概率可以分解为：

$$
P(X_1,\\ldots,X_n)=\\prod_{{i=1}}^n P(X_i|Pa(X_i))
$$

$Pa(X_i)$ 表示变量 $X_i$ 的父节点。公式表达的是：每个变量只需要依赖它的直接父节点，而不必直接依赖所有其他变量。

### 4.5 常见推理方式

贝叶斯网络常见推理包括：

1. 预测推理：已知原因，推断可能结果。
2. 诊断推理：已知结果，反推可能原因。
3. 解释消除：一个原因被确认后，其他可能原因的概率下降。
4. 敏感性分析：观察某个证据变化后，目标变量概率如何变化。

这些推理方式让贝叶斯网络适合处理不完整、有噪声、需要解释的证据。

### 4.6 适用场景与局限

贝叶斯网络适合医疗诊断、风险评估、故障定位、决策支持和知识推理。它的优势是可解释，能把领域知识和数据结合起来。

局限在于结构设计不容易，条件概率表也需要数据或专家知识支持。变量很多时，精确推理会变得复杂。对于图像、语音和大规模文本这类非结构化数据，贝叶斯网络通常需要与其他模型结合使用。

### 4.7 本节小结

贝叶斯网络把“常识关系”变成“可计算的概率关系”。它的重点不是画图好看，而是让复杂依赖可以被推理、更新和解释。

## 清单第 11 项：K-近邻（KNN）

### 5.1 一句话定位

K-近邻是一种实例学习方法。它几乎不训练模型，而是在预测时寻找最相似的 K 个样本，用邻居投票决定类别。

### 5.2 核心直觉

KNN 的想法接近日常判断：一个新对象和附近对象最像，就更可能属于附近对象的类别。它不急着总结一条公式化规则，而是保留训练样本，在需要判断时再比较距离。

![图 5-1：KNN 通过邻域投票判断新样本类别]({paths["knn"]}){{width=92%}}

这种方法很直观，也很容易解释。问题在于“附近”必须被量化，而量化方式就是距离或相似度。距离怎么定义，KNN 就会按怎样的标准理解相似。

### 5.3 一个容易理解的例子

假设要判断一个水果是苹果还是橙子，手里有两个特征：重量和颜色深浅。新水果落在特征空间中的某个位置后，KNN 会找到离它最近的几个已知水果。如果最近的 5 个样本里有 4 个是苹果，模型就把它判为苹果。

如果把 K 设得太小，结果容易被偶然噪声影响；如果 K 设得太大，局部差异会被平均掉。因此 K 的选择决定了模型边界是灵活还是平滑。

### 5.4 常见距离

欧氏距离适合连续数值特征：

$$
d(x_i,x_j)=\\sqrt{{\\sum_l (x_{{il}}-x_{{jl}})^2}}
$$

曼哈顿距离把各维差异直接相加：

$$
d(x_i,x_j)=\\sum_l |x_{{il}}-x_{{jl}}|
$$

文本向量和高维语义向量常用余弦相似度：

$$
sim(x_i,x_j)=\\frac{{x_i^Tx_j}}{{\\|x_i\\|\\|x_j\\|}}
$$

距离度量决定“相似”的含义。特征尺度也很重要，如果某个特征数值范围特别大，它会主导距离计算，所以 KNN 前通常要做标准化。

### 5.5 运行流程

1. 保存训练样本和标签。
2. 对新样本计算它与训练样本的距离。
3. 选出距离最近的 K 个样本。
4. 分类任务中做多数投票，回归任务中做平均或加权平均。
5. 输出预测结果。

KNN 的训练很轻，预测较重。数据越多，每次预测要比较的样本也越多。

### 5.6 适用场景与局限

KNN 适合小规模数据、低维特征、边界不规则但局部相似性明显的任务，也常被用作基线模型。它直观、无需复杂训练、预测结果容易解释。

局限包括预测速度慢、高维空间中距离变得不可靠、对无关特征敏感、对特征尺度敏感。数据量很大时，通常需要 KD-Tree、Ball Tree 或近似最近邻等加速方法。

### 5.7 本节小结

KNN 的核心是“近朱者赤”。它把分类问题转化为邻域投票问题，简单可靠，但依赖好的特征、合适的距离和合理的 K 值。

## 清单第 12 项：支持向量机（SVM）

### 6.1 一句话定位

支持向量机是一种最大间隔分类器。它不只追求把训练样本分开，还追求分类边界离两类最近样本都尽可能远。

### 6.2 核心直觉

如果有很多条直线都能把两类样本分开，哪一条更好？SVM 的答案是：选择中间余量最大的那一条。余量越大，新样本稍微有点扰动时越不容易被分错。

![图 6-1：SVM 寻找最大间隔边界，支持向量决定边界位置]({paths["svm"]}){{width=92%}}

被边界“贴得最近”的样本叫支持向量。它们数量可能不多，却决定了分类面的位置。远离边界的大量样本，对最终边界影响反而较小。

### 6.3 一个容易理解的例子

在文本分类中，邮件或新闻可以被表示为高维词向量。很多词可能没有用，真正有区分力的是少数关键词组合。SVM 在高维空间中寻找最大间隔边界，因此曾经广泛用于文本分类、图像手工特征分类和生物信息分类。

这个例子也说明，SVM 很适合“样本不算特别多，但特征维度很高”的任务。

### 6.4 必要公式

硬间隔 SVM 的目标是让边界尽可能宽：

$$
\\min_{{w,b}} \\frac{{1}}{{2}}\\|w\\|^2
$$

同时要求所有样本被正确分开：

$$
y_i(w^Tx_i+b)\\ge 1
$$

现实数据往往不能完全分开，所以软间隔 SVM 允许少量错误：

$$
\\min_{{w,b,\\xi}} \\frac{{1}}{{2}}\\|w\\|^2 + C\\sum_i \\xi_i
$$

$C$ 控制“尽量少犯错”和“保持大间隔”之间的折中。

### 6.5 核技巧

当原始空间中无法用直线分开样本时，SVM 可以通过核函数把问题转到更高维的特征空间。核函数写作：

$$
K(x_i,x_j)=\\phi(x_i)^T\\phi(x_j)
$$

直观理解是：模型不需要真的把样本坐标全部算到高维空间，只需要知道两个样本在高维空间中的内积。常见核函数包括线性核、多项式核、RBF 核和 Sigmoid 核，其中 RBF 核常用于非线性边界。

### 6.6 适用场景与局限

SVM 适合中小规模数据、高维特征和边界相对清晰的分类任务。它理论基础扎实，在样本数量有限时经常表现稳定。

局限是大规模训练成本高，核函数和参数选择依赖经验，概率输出不自然。面对海量图像、语音和文本数据时，深度学习模型通常更适合端到端学习。

### 6.7 本节小结

SVM 的核心是“分开还不够，要留足安全距离”。支持向量决定边界，核技巧让线性方法能够处理非线性问题。

## 7. 横向比较

| 序号 | 模型 | 类型 | 核心思想 | 适合问题 | 主要局限 |
|---|---|---|---|---|---|
| 7 | 感知机 | 判别式线性分类 | 分错就修正边界 | 线性二分类、神经网络入门 | 只能表达线性边界 |
| 8 | GMM | 生成式概率模型 | 多个概率团块混合产生数据 | 软聚类、密度估计、异常检测 | 需指定成分数，对初始化敏感 |
| 9 | HMM | 序列概率模型 | 隐藏状态按时间生成观测 | 语音、文本、动作、故障序列 | 难表达长期依赖 |
| 10 | 贝叶斯网络 | 概率图模型 | 用图表达变量条件依赖 | 医疗诊断、风险评估、知识推理 | 结构和概率表构造困难 |
| 11 | KNN | 实例学习 | 看最近邻居如何投票 | 小规模、低维、局部相似明显任务 | 预测慢，高维距离失效 |
| 12 | SVM | 最大间隔分类 | 找到最稳的分类边界 | 中小样本、高维分类 | 大规模训练与核选择困难 |

这六个模型并不是互相替代的关系，而是从不同角度回答模式识别问题。感知机和 SVM 关注边界，GMM、HMM、贝叶斯网络关注概率结构，KNN 关注样本相似性。理解这三类思路后，再学习决策树、神经网络、Transformer 或生成模型，会更容易抓住它们各自解决的问题。

## 8. 结语

模式识别课程中的经典模型并不只是历史内容。它们提供了几种稳定的思考方式：用边界做分类、用概率描述不确定性、用邻近样本判断相似性、用图结构表达依赖关系。现代 AI 系统规模更大、数据更多、模型更复杂，但很多基本问题仍然可以回到这些经典模型中找到解释。

对初学者来说，掌握这些模型的最好方式不是背公式，而是把每个公式和具体问题对应起来：感知机的公式在移动边界，GMM 的公式在混合分布，HMM 的公式在连接状态和观测，贝叶斯网络的公式在拆解联合概率，KNN 的距离在定义相似，SVM 的目标函数在扩大安全间隔。这样学习时会更清楚每个模型为什么存在、解决什么问题、适合用在什么地方。

## 参考资料

1. Richard O. Duda, Peter E. Hart, David G. Stork. *Pattern Classification*.
2. Christopher M. Bishop. *Pattern Recognition and Machine Learning*.
3. Tom M. Mitchell. *Machine Learning*.
4. Stuart Russell, Peter Norvig. *Artificial Intelligence: A Modern Approach*.
"""


def postprocess_docx() -> None:
    def set_style_font(style, name: str, size: float | None = None) -> None:
        style.font.name = name
        if size is not None:
            style.font.size = Pt(size)
        rpr = style._element.get_or_add_rPr()
        rfonts = rpr.rFonts
        if rfonts is not None:
            rfonts.set(qn("w:ascii"), name)
            rfonts.set(qn("w:hAnsi"), name)
            rfonts.set(qn("w:eastAsia"), name)

    def set_run_font(run, name: str, size: float | None = None) -> None:
        run.font.name = name
        if size is not None:
            run.font.size = Pt(size)
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.rFonts
        if rfonts is not None:
            rfonts.set(qn("w:ascii"), name)
            rfonts.set(qn("w:hAnsi"), name)
            rfonts.set(qn("w:eastAsia"), name)

    doc = Document(DOCX_PATH)
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    styles = doc.styles
    for style_name in ["Normal", "Body Text", "First Paragraph"]:
        if style_name in styles:
            style = styles[style_name]
            set_style_font(style, "PingFang SC", 10.5)
            style.paragraph_format.line_spacing = 1.25
            style.paragraph_format.space_after = Pt(6)

    for style_name, size in [("Title", 20), ("Heading 1", 18), ("Heading 2", 15), ("Heading 3", 12.5)]:
        if style_name in styles:
            style = styles[style_name]
            set_style_font(style, "PingFang SC", size)
            style.paragraph_format.space_before = Pt(10)
            style.paragraph_format.space_after = Pt(6)

    for para in doc.paragraphs:
        text = para.text.strip()
        if text.startswith("图 ") or text.startswith("图"):
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                set_run_font(run, "PingFang SC", 9)
        elif para.style.name in {"Heading 1", "Title"}:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for table in doc.tables:
        if "Table Grid" in styles:
            table.style = "Table Grid"
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    para.paragraph_format.space_after = Pt(2)
                    for run in para.runs:
                        set_run_font(run, "PingFang SC", 9)

    doc.save(DOCX_PATH)


def main() -> None:
    paths = {
        "overview": use_codex_generated("overview.png"),
        "perceptron": use_codex_generated("perceptron.png"),
        "gmm": use_codex_generated("gmm.png"),
        "hmm": use_codex_generated("hmm.png"),
        "bayes": use_codex_generated("bayes.png"),
        "knn": use_codex_generated("knn.png"),
        "svm": use_codex_generated("svm.png"),
    }
    MD_PATH.write_text(build_markdown(paths), encoding="utf-8")

    pandoc = shutil.which("pandoc")
    if not pandoc:
        raise SystemExit("pandoc not found")
    subprocess.run(
        [
            pandoc,
            str(MD_PATH),
            "-o",
            str(DOCX_PATH),
            "--from",
            "markdown+tex_math_dollars",
        ],
        cwd=ROOT,
        check=True,
    )
    postprocess_docx()


if __name__ == "__main__":
    main()
